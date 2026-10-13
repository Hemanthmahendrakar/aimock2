import io
from datetime import datetime, timezone

from app.extensions import db
from app.models import User, UserStatus
from app.services.candidate_profile_builder import build_candidate_profile


def _register_and_activate(client, *, institute_name, institute_email, admin_email, password):
    payload = {
        "institute_name": institute_name,
        "institute_email": institute_email,
        "admin_full_name": "Admin",
        "admin_email": admin_email,
        "password": password,
    }
    client.post("/api/v1/institute/auth/register", json=payload)
    user = db.session.query(User).filter(User.email == admin_email).first()
    user.status = UserStatus.ACTIVE
    user.email_verified_at = datetime.now(timezone.utc)
    db.session.commit()
    return user


def _login(client, email, password):
    return client.post(
        "/api/v1/institute/auth/login", json={"email": email, "password": password}
    )


def _setup_institute(client, suffix: str):
    _register_and_activate(
        client,
        institute_name=f"Institute {suffix}",
        institute_email=f"admissions@{suffix}.example",
        admin_email=f"admin@{suffix}.example",
        password="SuperSecret123",
    )
    _login(client, f"admin@{suffix}.example", "SuperSecret123")


def _create_student(client, email_suffix: str, name="Student"):
    resp = client.post(
        "/api/v1/institute/students",
        json={"name": name, "email": f"{email_suffix}@students.example"},
    )
    return resp.get_json()["data"]["id"]


# A minimal one-page PDF with "Hello Resume" as its content stream.
# The xref table below is intentionally not byte-exact (computing
# precise offsets by hand is error-prone) — pypdf recovers from a
# malformed/incomplete xref by scanning the file for objects directly,
# so this still parses. If a future pypdf version tightens that
# tolerance, regenerate this fixture with a real PDF library instead.
_MINIMAL_PDF = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>endobj
4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
5 0 obj<</Length 44>>
stream
BT /F1 12 Tf 10 100 Td (Hello Resume) Tj ET
endstream
endobj
xref
0 6
0000000000 65535 f 
trailer<</Size 6/Root 1 0 R>>
startxref
0
%%EOF"""


def _sample_docx_bytes() -> bytes:
    import docx

    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Education")
    document.add_paragraph("B.Tech Computer Science, XYZ University")
    document.add_paragraph("Skills")
    document.add_paragraph("Python, React, Docker")
    document.save(buf)
    return buf.getvalue()


class TestResumeUpload:
    def test_upload_pdf_success(self, client):
        _setup_institute(client, "rp1")
        student_id = _create_student(client, "rp1s1")

        resp = client.post(
            f"/api/v1/institute/students/{student_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "resume.pdf")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 201
        data = resp.get_json()["data"]
        assert data["version_number"] == 1
        # CELERY_TASK_ALWAYS_EAGER=True in tests, so processing already ran.
        assert data["extraction_status"] == "completed"

    def test_upload_docx_success(self, client):
        _setup_institute(client, "rp2")
        student_id = _create_student(client, "rp2s1")

        resp = client.post(
            f"/api/v1/institute/students/{student_id}/resume",
            data={"file": (io.BytesIO(_sample_docx_bytes()), "resume.docx")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 201
        data = resp.get_json()["data"]
        assert data["extraction_status"] == "completed"
        assert "Education" in data["extracted_text"]
        profile = data["candidate_profile"]
        assert profile["analysis_status"] == "completed"
        assert "Python" in profile["programming_languages"]

    def test_upload_rejects_unsupported_type(self, client):
        _setup_institute(client, "rp3")
        student_id = _create_student(client, "rp3s1")

        resp = client.post(
            f"/api/v1/institute/students/{student_id}/resume",
            data={"file": (io.BytesIO(b"just some text"), "resume.txt")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 422
        assert resp.get_json()["error"]["code"] == "invalid_file"

    def test_upload_for_nonexistent_student_404s(self, client):
        _setup_institute(client, "rp4")
        fake_id = "00000000-0000-0000-0000-000000000000"

        resp = client.post(
            f"/api/v1/institute/students/{fake_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "resume.pdf")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 404

    def test_second_upload_increments_version(self, client):
        _setup_institute(client, "rp5")
        student_id = _create_student(client, "rp5s1")

        client.post(
            f"/api/v1/institute/students/{student_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "v1.pdf")},
            content_type="multipart/form-data",
        )
        resp2 = client.post(
            f"/api/v1/institute/students/{student_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "v2.pdf")},
            content_type="multipart/form-data",
        )
        assert resp2.get_json()["data"]["version_number"] == 2

        current = client.get(f"/api/v1/institute/students/{student_id}/resume")
        assert current.get_json()["data"]["version_number"] == 2

        versions = client.get(f"/api/v1/institute/students/{student_id}/resume/versions")
        assert versions.get_json()["data"]["pagination"]["total"] == 2


class TestResumeRetrieval:
    def test_get_current_before_upload_404s(self, client):
        _setup_institute(client, "rr1")
        student_id = _create_student(client, "rr1s1")

        resp = client.get(f"/api/v1/institute/students/{student_id}/resume")
        assert resp.status_code == 404
        assert resp.get_json()["error"]["code"] == "resume_not_found"

    def test_get_specific_version(self, client):
        _setup_institute(client, "rr2")
        student_id = _create_student(client, "rr2s1")
        upload_resp = client.post(
            f"/api/v1/institute/students/{student_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "resume.pdf")},
            content_type="multipart/form-data",
        )
        version_id = upload_resp.get_json()["data"]["id"]

        resp = client.get(
            f"/api/v1/institute/students/{student_id}/resume/versions/{version_id}"
        )
        assert resp.status_code == 200
        assert resp.get_json()["data"]["id"] == version_id


class TestResumeTenantIsolation:
    def test_institute_b_cannot_upload_to_institute_a_student(self, client):
        _setup_institute(client, "tra")
        student_a_id = _create_student(client, "tras1")

        client.post("/api/v1/institute/auth/logout")
        _setup_institute(client, "trb")

        resp = client.post(
            f"/api/v1/institute/students/{student_a_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "resume.pdf")},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 404

    def test_institute_b_cannot_read_institute_a_resume(self, client):
        _setup_institute(client, "trc")
        student_a_id = _create_student(client, "trcs1")
        client.post(
            f"/api/v1/institute/students/{student_a_id}/resume",
            data={"file": (io.BytesIO(_MINIMAL_PDF), "resume.pdf")},
            content_type="multipart/form-data",
        )

        client.post("/api/v1/institute/auth/logout")
        _setup_institute(client, "trd")

        resp = client.get(f"/api/v1/institute/students/{student_a_id}/resume")
        assert resp.status_code == 404


class TestCandidateProfileBuilder:
    """Exercises the heuristic parser directly — faster and more
    reliable than round-tripping through file bytes for every case."""

    def test_extracts_sections(self):
        text = (
            "John Doe\n"
            "Education\n"
            "B.Tech Computer Science, XYZ University, 2026\n"
            "Projects\n"
            "Built a mock interview platform\n"
            "Skills\n"
            "Python, React, Docker\n"
            "Certifications\n"
            "AWS Certified Developer\n"
        )
        profile = build_candidate_profile(text)
        assert "B.Tech Computer Science, XYZ University, 2026" in profile["education"]
        assert "Built a mock interview platform" in profile["projects"]
        assert "AWS Certified Developer" in profile["certifications"]
        assert "Python" in profile["skills"]
        assert "React" in profile["skills"]

    def test_detects_languages_and_technologies(self):
        text = "Experienced with Python, Java, and deployed services using Docker and AWS."
        profile = build_candidate_profile(text)
        assert "Python" in profile["programming_languages"]
        assert "Java" in profile["programming_languages"]
        assert "Docker" in profile["technologies"]
        assert "AWS" in profile["technologies"]
        # Not mentioned, should not appear.
        assert "Kotlin" not in profile["programming_languages"]

    def test_missing_sections_return_empty_lists(self):
        profile = build_candidate_profile("Just some random text with no headers.")
        assert profile["education"] == []
        assert profile["certifications"] == []
